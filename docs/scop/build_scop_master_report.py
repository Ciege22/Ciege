"""
SCOP Master Report Generator
==============================
Generates the 9-tab Excel Master Report (Summary, Outstanding List, Pathwave
Pending/Complete/OAD Sites, QuickBase Pending/Complete, Fully Complete, Filter
Cheat Sheet). This is the Python/openpyxl counterpart to scop_calculations.js —
if this codebase's Excel export layer is Python-based, use this directly; if
it's Node-based, port the structure using a library like exceljs, keeping the
same tab layout, column sets, and the two bug-fix notes below.

Expects a pandas DataFrame already scoped to Nokia and read with the header
row auto-detected (see the header-row hazard in SCOP_View_Build_Spec.md).

Usage:
    df = pd.read_excel(path, sheet_name="HOPs", header=detected_header_row)
    df = df[df["Service provider"].astype(str).str.strip().str.lower() == "nokia"]
    build_master_report(df, "SCOP_Master_Report.xlsx", as_of_date=pd.Timestamp.now())
"""

import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.page import PageMargins

# ---------------------------------------------------------------------------
# Styling constants
# ---------------------------------------------------------------------------
NAVY = "124191"
OAD_COLOR = "2C3E70"
WHITE = "FFFFFF"
RED_FILL = "FDEDEC"
TEAL_FILL = "E6F7F9"
OAD_FILL = "EEF2FA"
RED_TXT = "C0392B"
AMBER_TXT = "B7791F"
GREEN_TXT = "1E8449"

THIN = Side(style="thin", color="D0D0D0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SIMPLE_HEADERS = ["HOP", "Path ID", "GC", "CM", "Near Site A", "Far Site B"]

# ---------------------------------------------------------------------------
# Item ownership — hard-coded, see scop_calculations.js for why this can't be
# re-derived from the tracker file itself (no ownership metadata in the live file).
# ---------------------------------------------------------------------------
PATHWAVE_GC_ITEMS = [
    ("Install Photos", "Viaero Install Photos Status Site A", "Viaero Install Photo Status B"),
    ("Decom Photos NQR", "Site A Decom Photos NQR", "Site B Decom Photos NQR"),
    ("Install Photos NQR", "Site A Install Photos NQR", "Site B Install Photos NQR"),
    ("Red-Line CD", "Site A Red-Line CD", "Site B Red Line CDs"),
    ("Decom Asset Form", "Site A Decom Asset Form", "Site B Decom Asset Form"),
    ("POD", "Site A POD", "Site B POD"),
]
QUICKBASE_ALL_ITEMS = [  # ALL Nokia-owned — used for internal tracking only, never GC-facing
    ("B2B Test", "QB Site A B2B Test", "Site B B2B Test"),
    ("RFC2455 Test", "QB Site A RFC2455 Test", "QB Site B RFC2455 Test"),
    ("BER Report", "QB Site A BER Report", "QB Site B BER Report"),
    ("Consolidated Asset Form", "QB Site A Consolidated Asset Form", "QB Site B Consolidated Asset Form"),
    ("As-Built Final CD", "QB Site A As-Built Final CD", "QB Site B As-Built Final CD"),
    ("As Built WP", "QB Site A As Built WP", "QB Site B As Built WP"),
    ("As-Built LLD", "QB Site A As-Built LLD", "QB Site B As-Built LLD"),
    ("KPI Report", "QB Site A KPI Report", "QB Site A KPI Report3"),  # sic — mislabeled in source, this IS Site B's field
    ("POD", "Site A POD2", "Site B POD4"),
    ("HW POD", "HW POD", "HW POD2"),
]

GC_ALIAS_MAP = {"wavelink": "WaveLink", "viking/capital tower": "Viking"}  # reuse the real setting instead of this constant


# ---------------------------------------------------------------------------
# Normalization helpers
# ---------------------------------------------------------------------------
def norm(v):
    if pd.isna(v):
        return None
    s = str(v).strip()
    return s.lower() if s else None


def is_done(v):
    s = norm(v)
    return s is not None and s in ("yes", "accepted", "complete")


def status_bucket(v):
    s = norm(v)
    return "Complete" if (s is not None and s.startswith("complete")) else "Pending"


def norm_gc(v, alias_map=GC_ALIAS_MAP):
    if pd.isna(v):
        return None
    s = str(v).strip()
    if s == "":
        return None
    return alias_map.get(s.lower(), s)


def aging_color(days):
    if days is None:
        return "808080"
    if days >= 60:
        return RED_TXT
    if days >= 30:
        return AMBER_TXT
    return GREEN_TXT


def prepare_dataset(df, as_of_date, gc_alias_map=GC_ALIAS_MAP):
    """
    Adds all derived columns to a Nokia-scoped DataFrame. Call this once,
    then pass the result into build_master_report().

    CRITICAL: `Construction Complete Actual` blank cells can parse as a
    time-only artifact (e.g. datetime.time(0,0)) rather than a true null,
    depending on how the sheet was read. ALWAYS use pd.to_datetime(...,
    errors="coerce") and check notna() on the PARSED result — never a raw
    notna() check on the unparsed column. This exact bug silently inflated
    a "construction complete" count to 100% once already.
    """
    df = df.copy()
    df["constr_date"] = pd.to_datetime(df["Construction Complete Actual"], errors="coerce")
    df["constr_complete"] = df["constr_date"].notna()
    df["aging_days"] = (as_of_date - df["constr_date"]).dt.days
    df["GC"] = df["General Contractor"].apply(lambda v: norm_gc(v, gc_alias_map))
    df["is_oad"] = df["One and Done"].astype(str).str.contains("oad", case=False, na=False)
    df["pw_bucket"] = df["One and Done"].apply(status_bucket)
    df["qb_bucket"] = df["Nokia Quickbase Deliverable Status"].apply(status_bucket)
    return df


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------
def _write_simple_list(ws, sub_df, title):
    r = 1
    ws.merge_cells(f"A{r}:F{r}")
    ws.cell(row=r, column=1, value=title).font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    r += 2
    for col_i, h in enumerate(SIMPLE_HEADERS, start=1):
        c = ws.cell(row=r, column=col_i, value=h)
        c.font = Font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    r += 1
    for _, row in sub_df.iterrows():
        for col_i, h in enumerate(SIMPLE_HEADERS, start=1):
            c = ws.cell(row=r, column=col_i, value=row[h])
            c.border = BORDER
            c.font = Font(size=10)
        r += 1
    for col_i, h in enumerate(SIMPLE_HEADERS, start=1):
        ws.column_dimensions[get_column_letter(col_i)].width = max(14, len(h) + 4)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws.print_area = f"A1:F{r - 1}"


def _build_itemized_rows(source_df, item_pairs):
    out = []
    for _, row in source_df.iterrows():
        for site, site_letter in [("A", "A"), ("B", "B")]:
            missing = []
            for label, col_a, col_b in item_pairs:
                col = col_a if site == "A" else col_b
                if not is_done(row[col]):
                    missing.append(label)
            if not missing:
                continue
            out.append({
                "HOP": row["HOP"], "Path ID": row.get("Path ID"), "GC": row["GC"], "CM": row.get("Site CM"),
                "Near Site A": row.get("Near Site Name A"), "Far Site B": row.get("Far Site Name B"),
                "Site": f"Site {site_letter}", "Aging": int(row["aging_days"]) if pd.notna(row["aging_days"]) else None,
                "Missing Items": ", ".join(missing),
            })
    return out


def _write_itemized_table(ws, rows, title, fill_color, header_color=NAVY):
    r = 1
    ws.merge_cells(f"A{r}:I{r}")
    ws.cell(row=r, column=1, value=title).font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    r += 1
    ws.merge_cells(f"A{r}:I{r}")
    ws.cell(row=r, column=1, value="Sorted oldest-first (highest days since construction complete) — clean up the top rows first").font = \
        Font(name="Calibri", size=10, italic=True, color="6B7280")
    r += 2
    headers = ["HOP", "Path ID", "GC", "CM", "Near Site A", "Far Site B", "Site", "Days Since Complete", "Missing Items"]
    for col_i, h in enumerate(headers, start=1):
        c = ws.cell(row=r, column=col_i, value=h)
        c.font = Font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=header_color)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    header_row = r
    r += 1
    rows_sorted = sorted(rows, key=lambda x: -(x["Aging"] if x["Aging"] is not None else 0))
    for row_data in rows_sorted:
        ws.cell(row=r, column=1, value=row_data["HOP"]).border = BORDER
        ws.cell(row=r, column=2, value=row_data["Path ID"]).border = BORDER
        ws.cell(row=r, column=3, value=row_data["GC"]).border = BORDER
        ws.cell(row=r, column=4, value=row_data["CM"]).border = BORDER
        ws.cell(row=r, column=5, value=row_data["Near Site A"]).border = BORDER
        ws.cell(row=r, column=6, value=row_data["Far Site B"]).border = BORDER
        sc = ws.cell(row=r, column=7, value=row_data["Site"]); sc.border = BORDER; sc.alignment = Alignment(horizontal="center")
        ac = ws.cell(row=r, column=8, value=row_data["Aging"]); ac.border = BORDER; ac.alignment = Alignment(horizontal="center")
        ac.font = Font(bold=True, color=aging_color(row_data["Aging"]))
        mc = ws.cell(row=r, column=9, value=row_data["Missing Items"]); mc.border = BORDER
        mc.fill = PatternFill("solid", fgColor=fill_color)
        mc.alignment = Alignment(wrap_text=True, vertical="center")
        for col_i in range(1, 7):
            ws.cell(row=r, column=col_i).font = Font(size=10)
        r += 1
    widths = [20, 10, 13, 9, 15, 15, 6, 10, 32]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{header_row + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)
    ws.print_area = f"A1:I{r - 1}"


def _write_oad_table(ws, source_df):
    r = 1
    ws.merge_cells("A1:G1")
    ws.cell(row=r, column=1, value="Pathwave — OAD Sites (Awaiting OAD, Not a GC Checklist Item)").font = \
        Font(name="Calibri", size=14, bold=True, color=NAVY)
    r += 1
    ws.merge_cells("A2:G2")
    ws.cell(row=r, column=1, value="Sorted oldest-first. These sites are blocked on OAD, not on GC-submitted checklist items.").font = \
        Font(name="Calibri", size=10, italic=True, color="6B7280")
    r += 2
    headers = ["HOP", "Path ID", "GC", "CM", "Near Site A", "Far Site B", "Days Since Complete", "Note"]
    for col_i, h in enumerate(headers, start=1):
        c = ws.cell(row=r, column=col_i, value=h)
        c.font = Font(bold=True, color=WHITE)
        c.fill = PatternFill("solid", fgColor=OAD_COLOR)
        c.alignment = Alignment(horizontal="center")
        c.border = BORDER
    header_row = r
    r += 1
    for _, row in source_df.sort_values("aging_days", ascending=False).iterrows():
        ws.cell(row=r, column=1, value=row["HOP"]).border = BORDER
        ws.cell(row=r, column=2, value=row.get("Path ID")).border = BORDER
        ws.cell(row=r, column=3, value=row["GC"]).border = BORDER
        ws.cell(row=r, column=4, value=row.get("Site CM")).border = BORDER
        ws.cell(row=r, column=5, value=row.get("Near Site Name A")).border = BORDER
        ws.cell(row=r, column=6, value=row.get("Far Site Name B")).border = BORDER
        days = int(row["aging_days"]) if pd.notna(row["aging_days"]) else None
        ac = ws.cell(row=r, column=7, value=days); ac.border = BORDER; ac.alignment = Alignment(horizontal="center")
        ac.font = Font(bold=True, color=aging_color(days))
        # NOTE: this cell holds free text from the tracker (e.g. "pending OAD",
        # "OAD ETA 9/11") — NEVER prefix it with "=" or any formula-like
        # character. A prior version accidentally wrote a value starting with
        # "=" and Excel/openpyxl silently evaluated it as a formula, producing
        # Err:509 instead of the intended text.
        nc = ws.cell(row=r, column=8, value=row["One and Done"]); nc.border = BORDER
        nc.fill = PatternFill("solid", fgColor=OAD_FILL)
        for col_i in range(1, 7):
            ws.cell(row=r, column=col_i).font = Font(size=10)
        r += 1
    widths = [20, 10, 13, 9, 15, 15, 10, 38]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = f"A{header_row + 1}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins = PageMargins(left=0.3, right=0.3, top=0.4, bottom=0.4)
    ws.print_area = f"A1:H{r - 1}"


# ---------------------------------------------------------------------------
# Main entry point
# ---------------------------------------------------------------------------
def build_master_report(df, output_path, as_of_date=None, gc_alias_map=GC_ALIAS_MAP):
    """
    df: pandas DataFrame already read from the HOPs tab with the correct header
        row, already filtered to Service provider == "Nokia" (case-insensitive).
    output_path: where to write the .xlsx file.
    as_of_date: pd.Timestamp for aging calculations. Defaults to now.
    """
    if as_of_date is None:
        as_of_date = pd.Timestamp.now()

    df = prepare_dataset(df, as_of_date, gc_alias_map)
    tracked = df[df["constr_complete"]].copy()

    n_oad = int(tracked["is_oad"].sum())
    n_pw_pending_true = int(((tracked["pw_bucket"] == "Pending") & (~tracked["is_oad"])).sum())
    n_pw_complete = int((tracked["pw_bucket"] == "Complete").sum())
    n_qb_complete = int((tracked["qb_bucket"] == "Complete").sum())
    n_qb_pending = int((tracked["qb_bucket"] == "Pending").sum())
    fully_complete = int(((tracked["pw_bucket"] == "Complete") & (tracked["qb_bucket"] == "Complete")).sum())

    wb = openpyxl.Workbook()

    # --- Summary ---
    ws = wb.active
    ws.title = "Summary"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:D1")
    ws["A1"] = "SCOP Master Report — Summary"
    ws["A1"].font = Font(name="Calibri", size=18, bold=True, color=NAVY)
    ws.merge_cells("A2:D2")
    ws["A2"] = f"Generated {as_of_date.strftime('%m/%d/%Y')}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="6B7280")

    summary_rows = [
        ("Total Nokia HOPs", len(df)),
        ("Construction Complete (tracked)", int(df["constr_complete"].sum())),
        ("Outstanding (not yet construction complete)", int((~df["constr_complete"]).sum())),
        ("Pathwave — Complete", n_pw_complete),
        ("Pathwave — In Progress (true GC action, excl. OAD)", n_pw_pending_true),
        ("Pathwave — OAD (tracked separately)", n_oad),
        ("QuickBase — Complete", n_qb_complete),
        ("QuickBase — Pending", n_qb_pending),
        ("Fully Complete (both Pathwave + QuickBase)", fully_complete),
    ]
    r = 4
    ws.cell(row=r, column=1, value="Metric").font = Font(bold=True, color=WHITE)
    ws.cell(row=r, column=2, value="Count").font = Font(bold=True, color=WHITE)
    ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor=NAVY)
    r += 1
    for label, val in summary_rows:
        ws.cell(row=r, column=1, value=label).font = Font(size=11)
        ws.cell(row=r, column=2, value=val).font = Font(size=11, bold=True)
        r += 1
    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 14

    # --- Outstanding List ---
    ws2 = wb.create_sheet("Outstanding List")
    ws2.sheet_view.showGridLines = False
    sub = df[~df["constr_complete"]][["HOP", "Path ID", "GC", "Site CM", "Near Site Name A", "Far Site Name B"]].rename(
        columns={"Site CM": "CM", "Near Site Name A": "Near Site A", "Far Site Name B": "Far Site B"})
    _write_simple_list(ws2, sub, "Outstanding — Not Yet Construction Complete")

    # --- Pathwave Pending (itemized, OAD excluded) ---
    ws3 = wb.create_sheet("Pathwave Pending")
    ws3.sheet_view.showGridLines = False
    pw_rows_df = tracked[(tracked["pw_bucket"] == "Pending") & (~tracked["is_oad"])]
    pw_items = _build_itemized_rows(pw_rows_df, PATHWAVE_GC_ITEMS)
    _write_itemized_table(ws3, pw_items, "Pathwave — Pending Items (GC-Owned, Action Needed)", RED_FILL)

    # --- Pathwave Complete ---
    ws4 = wb.create_sheet("Pathwave Complete")
    ws4.sheet_view.showGridLines = False
    sub4 = tracked[tracked["pw_bucket"] == "Complete"][["HOP", "Path ID", "GC", "Site CM", "Near Site Name A", "Far Site Name B"]].rename(
        columns={"Site CM": "CM", "Near Site Name A": "Near Site A", "Far Site Name B": "Far Site B"})
    _write_simple_list(ws4, sub4, "Pathwave — Complete (Pending Viaero Approval)")

    # --- Pathwave OAD Sites ---
    ws_oad = wb.create_sheet("Pathwave OAD Sites")
    ws_oad.sheet_view.showGridLines = False
    _write_oad_table(ws_oad, tracked[tracked["is_oad"]])

    # --- QuickBase Pending (Nokia-internal, NOT GC-facing) ---
    ws5 = wb.create_sheet("QuickBase Pending")
    ws5.sheet_view.showGridLines = False
    qb_rows_df = tracked[tracked["qb_bucket"] == "Pending"]
    qb_items = _build_itemized_rows(qb_rows_df, QUICKBASE_ALL_ITEMS)
    _write_itemized_table(ws5, qb_items, "QuickBase — Pending Items (Nokia-Internal, NOT a GC Action Item)", TEAL_FILL)
    ws5["A1"] = "QuickBase — Pending Items (Nokia-Internal, NOT a GC Action Item)"
    ws5["A2"] = "All QuickBase items are Nokia-owned. This list is for internal engineering tracking only — never send to a GC. Sorted oldest-first."

    # --- QuickBase Complete ---
    ws6 = wb.create_sheet("QuickBase Complete")
    ws6.sheet_view.showGridLines = False
    sub6 = tracked[tracked["qb_bucket"] == "Complete"][["HOP", "Path ID", "GC", "Site CM", "Near Site Name A", "Far Site Name B"]].rename(
        columns={"Site CM": "CM", "Near Site Name A": "Near Site A", "Far Site Name B": "Far Site B"})
    _write_simple_list(ws6, sub6, "QuickBase — Complete")

    # --- Fully Complete ---
    ws7 = wb.create_sheet("Fully Complete")
    ws7.sheet_view.showGridLines = False
    sub7 = tracked[(tracked["pw_bucket"] == "Complete") & (tracked["qb_bucket"] == "Complete")][
        ["HOP", "Path ID", "GC", "Site CM", "Near Site Name A", "Far Site Name B"]
    ].rename(columns={"Site CM": "CM", "Near Site Name A": "Near Site A", "Far Site Name B": "Far Site B"})
    _write_simple_list(ws7, sub7, "Fully Complete — Both Pathwave & QuickBase Done (Ready for Viaero Approval)")

    # --- Filter Cheat Sheet ---
    ws8 = wb.create_sheet("Filter Cheat Sheet")
    ws8.sheet_view.showGridLines = False
    ws8.merge_cells("A1:C1")
    ws8["A1"] = "Filter Cheat Sheet — How to Rebuild These Lists Manually"
    ws8["A1"].font = Font(name="Calibri", size=16, bold=True, color=NAVY)
    cheat_rows = [
        ("List", "Filter On", "Condition"),
        ("Scope (which HOPs are yours)", 'Column "Service provider" (position may shift — filter by column NAME, not letter)',
         'Exact value "Nokia" (case-insensitive)'),
        ("Outstanding (not built yet)", "Column P — Construction Complete Actual", "Blank / no valid date"),
        ("Construction-complete (tracked universe)", "Column P — Construction Complete Actual", "Has a valid date"),
        ("Pathwave — OAD (separate from Pending)", "Column S — One and Done",
         'Contains "OAD" anywhere, case-insensitive AND Column P has a date. Check BEFORE "Pending" below.'),
        ("Pathwave — Pending (true GC action)", "Column S — One and Done",
         'Does NOT start with "Complete" AND does NOT contain "OAD" AND Column P has a date'),
        ("Pathwave — Complete", "Column S — One and Done", 'Starts with "Complete" AND Column P has a date'),
        ("QuickBase — Pending (Nokia-internal only)", "Column BT — Nokia Quickbase Deliverable Status",
         'Does NOT start with "Complete" AND Column P has a date. ALL QuickBase items are Nokia-owned — never GC-facing.'),
        ("QuickBase — Complete", "Column BT — Nokia Quickbase Deliverable Status", 'Starts with "Complete" AND Column P has a date'),
        ("Fully Complete (approvals-ready)", "Columns S + BT together", "Both start with Complete AND Column P has a date"),
        ("Pathwave GC action items (weekly report)", "Hard-coded item list",
         "6 of 8 Pathwave items: excludes Asset Form & Packing Slip (Nokia-owned)."),
        ("QuickBase GC action items", "N/A", "NONE. All QuickBase items are Nokia-owned. Never appears on a GC report or email."),
    ]
    r = 3
    for i, row in enumerate(cheat_rows):
        for col_i, val in enumerate(row, start=1):
            c = ws8.cell(row=r, column=col_i, value=val)
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if i == 0:
                c.font = Font(bold=True, color=WHITE)
                c.fill = PatternFill("solid", fgColor=NAVY)
            else:
                c.font = Font(size=10.5)
        ws8.row_dimensions[r].height = 40 if i > 0 else 20
        r += 1
    ws8.column_dimensions["A"].width = 32
    ws8.column_dimensions["B"].width = 38
    ws8.column_dimensions["C"].width = 54
    ws8.page_setup.orientation = "landscape"
    ws8.page_setup.fitToWidth = 1
    ws8.page_setup.fitToHeight = 0
    ws8.sheet_properties.pageSetUpPr.fitToPage = True
    ws8.page_margins = PageMargins(left=0.4, right=0.4, top=0.5, bottom=0.5)
    ws8.print_area = f"A1:C{r - 1}"

    wb.save(output_path)
    return output_path
