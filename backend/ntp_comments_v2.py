"""NTP comments for Deck Builder V2.

The NTP tab on the V2 page edits comments in Supabase instead of an uploaded
workbook. Every build reads the saved comments back into the NTP slides, and
the Excel export is only for sending to the customer after the call.

Store shape (one pm_updates_cache row, id STORE_ID):
    { "<Mon> Pending NTP": { "<HOP>": {"comment": str, "status": str} } }
"""
import io
import json
import threading
from datetime import datetime

import openpyxl
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STORE_ID = 'ntp-comments-v2'
# One writer at a time: every save is read-modify-write on the same row, so
# overlapping saves could otherwise overwrite each other.
STORE_LOCK = threading.Lock()
COMMENT_HEADER = 'COMMENT (fill after call)'
STATUS_OPTIONS = ['Pending', 'In Progress', 'Action Taken', 'Needs Attention']
SECTIONS = [
    ('External', '🔴 EXTERNAL BLOCKERS'),
    ('Other', '🟠 OTHER'),
    ('Program Team', '🔵 PROGRAM TEAM ACTIONS'),
]
HEADERS = ['HOP', 'Path ID', 'Category', 'Action Owner', 'GC', 'FC Start', 'FC End',
           'NTP Blocker / Waiting On', COMMENT_HEADER, 'STATUS']
COL_WIDTHS = [45, 14, 14, 18, 14, 10, 10, 50, 58, 16]


def sheet_key(month_abbr: str) -> str:
    return f'{month_abbr} Pending NTP'


def load_store(sb) -> dict:
    if not sb:
        return {}
    try:
        res = sb.table('pm_updates_cache').select('updates').eq('id', STORE_ID).execute()
        if res.data:
            return json.loads(res.data[0]['updates'])
    except Exception as e:
        print(f'[ntp-v2] load failed: {e}', flush=True)
    return {}


def save_store(sb, store: dict):
    # Errors propagate on purpose: a save that silently fails is how comments
    # disappear (see the SPO pricing bug), so the caller must report it.
    sb.table('pm_updates_cache').upsert({
        'id': STORE_ID,
        'updates': json.dumps(store),
        'updated_at': datetime.utcnow().isoformat(),
    }).execute()


def comments_by_sheet(store: dict) -> dict:
    """{sheet_key: {HOP: comment}} — the shape the NTP slide code already reads."""
    out = {}
    for sheet, hops in store.items():
        out[sheet] = {hop: v.get('comment', '') for hop, v in hops.items() if v.get('comment')}
    return out


def set_entry(store: dict, sheet: str, hop: str, comment=None, status=None) -> dict:
    entry = store.setdefault(sheet, {}).setdefault(hop, {'comment': '', 'status': 'Pending'})
    if comment is not None:
        entry['comment'] = comment
    if status is not None:
        entry['status'] = status
    entry['updated_at'] = datetime.utcnow().isoformat()
    return store


def import_workbook(path: str) -> dict:
    """Reads the comments and statuses from an NTP workbook (the Excel the
    team has been filling in) into the store shape. Only the month sheets
    are read; the Comments History sheet repeats the same text."""
    wb = openpyxl.load_workbook(path, data_only=True)
    store = {}
    for ws in wb.worksheets:
        if 'Pending NTP' not in ws.title:
            continue
        rows = list(ws.iter_rows(values_only=True))
        header_idx = next((i for i, r in enumerate(rows) if r and r[0] == 'HOP'), None)
        if header_idx is None:
            continue
        hdr = [str(c).strip() if c is not None else '' for c in rows[header_idx]]
        c_col = hdr.index(COMMENT_HEADER) if COMMENT_HEADER in hdr else None
        s_col = hdr.index('STATUS') if 'STATUS' in hdr else None
        k_col = hdr.index('Category') if 'Category' in hdr else None
        for r in rows[header_idx + 1:]:
            hop = str(r[0]).strip() if r and r[0] else ''
            if not hop or k_col is None or r[k_col] not in ('External', 'Other', 'Program Team'):
                continue
            comment = str(r[c_col]).strip() if c_col is not None and r[c_col] else ''
            status = str(r[s_col]).strip() if s_col is not None and r[s_col] else ''
            if comment or status:
                set_entry(store, ws.title, hop, comment=comment or '', status=status or 'Pending')
    return store


def apply_batch(sb, sheet: str, entries: list) -> dict:
    """Saves several HOP edits in one read-modify-write under the lock, then
    reads the row back and checks every edit is there before reporting success."""
    with STORE_LOCK:
        store = load_store(sb)
        for e in entries:
            set_entry(store, sheet, e['hop'], comment=e.get('comment'), status=e.get('status'))
        save_store(sb, store)
        landed = load_store(sb)
    missing = []
    for e in entries:
        got = landed.get(sheet, {}).get(e['hop'], {})
        if 'comment' in e and got.get('comment', '') != (e['comment'] or ''):
            missing.append(e['hop'])
        if 'status' in e and got.get('status') != e['status']:
            missing.append(e['hop'])
    if missing:
        raise RuntimeError(f'Save did not persist for: {", ".join(missing)}')
    return {'saved': len(entries), 'saved_at': datetime.utcnow().isoformat()}


def merge_store(base: dict, incoming: dict) -> dict:
    for sheet, hops in incoming.items():
        for hop, v in hops.items():
            set_entry(base, sheet, hop, comment=v.get('comment', ''), status=v.get('status', 'Pending'))
    return base


def _fmt_date(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ''
    try:
        ts = pd.Timestamp(v)
        return '' if pd.isna(ts) else ts.strftime('%m/%d')
    except Exception:
        return ''


def month_rows(pending_rows: list, store_sheet: dict) -> list:
    """Flattens one month's pending HOPs from compute_por_window into the rows
    the NTP tab and the export both use, with saved comment and status."""
    out = []
    for r in pending_rows:
        saved = store_sheet.get(r['HOP'], {})
        out.append({
            'hop': r['HOP'],
            'path_id': r.get('_path_id', ''),
            'category': r.get('cat', 'Other'),
            'owner': str(r.get('_ntp_owner', '') or ''),
            'gc': str(r.get('General Contractor', '') or ''),
            'fc_start': _fmt_date(r.get('MS15 Implementation Start F')),
            'fc_end': _fmt_date(r.get('MS16 Implementation Ends F')),
            'blocker': str(r.get('_ntp_wait', '') or ''),
            'comment': saved.get('comment', ''),
            'status': saved.get('status', 'Pending'),
        })
    order = {'External': 0, 'Other': 1, 'Program Team': 2}
    out.sort(key=lambda x: (order.get(x['category'], 1), x['fc_end'] or '99/99'))
    return out


def export_workbook(months: list) -> bytes:
    """months: [(sheet_name, rows)] with rows from month_rows(). Same layout
    as the workbook the team already uses, so the customer copy matches it."""
    NAVY, LT_BLUE = '124191', 'EEF2F7'
    CAT_FILLS = {'External': ('FFE5E5', 'C00000'), 'Other': ('FFF3E0', 'C55A11'),
                 'Program Team': ('E8F0FC', '124191')}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for sheet, rows in months:
        ws = wb.create_sheet(sheet)
        for ci, (hdr, cw) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
            c = ws.cell(row=1, column=ci, value=hdr)
            c.fill = PatternFill('solid', fgColor=NAVY)
            c.font = Font(name='Calibri', bold=True, size=10, color='FFFFFF')
            c.alignment = Alignment(horizontal='center', vertical='center')
            ws.column_dimensions[get_column_letter(ci)].width = cw
        ws.row_dimensions[1].height = 21.95
        ws.merge_cells('A2:J2')
        ws.cell(row=2, column=1, value='Sorted to match slides: External first → FC End oldest to newest.').fill = PatternFill('solid', fgColor=LT_BLUE)
        r_i = 3
        for cat, label in SECTIONS:
            group = [r for r in rows if r['category'] == cat]
            if not group:
                continue
            ws.merge_cells(start_row=r_i, start_column=1, end_row=r_i, end_column=10)
            ws.cell(row=r_i, column=1, value=label).font = Font(name='Calibri', bold=True, size=10)
            r_i += 1
            for r in group:
                vals = [r['hop'], r['path_id'], r['category'], r['owner'], r['gc'],
                        r['fc_start'], r['fc_end'], r['blocker'], r['comment'], r['status']]
                fill, fc = CAT_FILLS.get(cat, ('FFFFFF', '000000'))
                for ci, v in enumerate(vals, 1):
                    c = ws.cell(row=r_i, column=ci, value=v)
                    c.font = Font(name='Calibri', size=9, color=fc if ci == 3 else '000000')
                    c.fill = PatternFill('solid', fgColor=fill)
                    c.alignment = Alignment(wrap_text=True, vertical='top')
                r_i += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
