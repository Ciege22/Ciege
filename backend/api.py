import io
import os
import sys
import json
import base64
import shutil
import zipfile
import tempfile
import traceback
import logging
from datetime import datetime, timedelta

import requests
from flask import Flask, request, send_file, jsonify
from flask_cors import CORS
from werkzeug.utils import secure_filename
from supabase import create_client
from pptx import Presentation

# Ensure backend package path is importable when running this file directly
sys.path.append(os.path.dirname(__file__))
import build_deck
import build_deck_v2
import gr_tracker
import ntp_comments_v2

app = Flask(__name__)
CORS(app)
# Flask's default logger level is WARNING (since debug=False below), which
# would silently drop app.logger.info(...) calls — bump it to INFO so the
# ai_assistant diagnostic logging below actually reaches Railway's log
# stream instead of being swallowed.
app.logger.setLevel(logging.INFO)

SUPABASE_URL = os.environ.get('SUPABASE_URL', '')
SUPABASE_KEY = os.environ.get('SUPABASE_KEY', '')
supabase_client = create_client(SUPABASE_URL, SUPABASE_KEY) if SUPABASE_URL and SUPABASE_KEY else None

ANTHROPIC_API_KEY = os.environ.get('ANTHROPIC_API_KEY', '')
ANTHROPIC_MODEL = 'claude-sonnet-5'
AI_ASSISTANT_SYSTEM_PROMPT = (
	"You are Ciege AI, an intelligent program management assistant embedded in the Ciege platform "
	"for CJ, a Nokia Program Manager on the Viaero Wireless MW Construction Program (DON 444). "
	"You have access to real-time program data provided below. Answer questions concisely and "
	"accurately based on this data. When listing sites, format them clearly. When asked about "
	"blockers, check NTP status, material status, and vendor conflicts. When asked about a specific "
	"site, pull all available data for that HOP. Keep answers brief and actionable — CJ is often "
	"on a live call. Never make up data not in the context. If you don't have the data to answer, "
	"say so clearly."
)


def _save_uploaded_file(uploaded, dest_dir, field_name):
	if not uploaded:
		return ''
	filename = secure_filename(uploaded.filename) or field_name
	out_path = os.path.join(dest_dir, filename)
	uploaded.save(out_path)
	return out_path


@app.route('/build', methods=['POST'])
def build_endpoint():
	tmpdir = tempfile.mkdtemp(prefix='ciege_build_')
	try:
		# Read uploaded files
		tracker = request.files.get('tracker')
		previous_deck = request.files.get('previous_deck')
		snapshot = request.files.get('snapshot')
		ntp_comments = request.files.get('ntp_comments')
		deck_date = request.form.get('deck_date') or request.args.get('deck_date')

		if not previous_deck or not deck_date:
			return jsonify({'error': 'Missing required fields: previous_deck, deck_date'}), 400

		tracker_path = _save_uploaded_file(tracker, tmpdir, 'tracker.xlsx') if tracker else ''
		previous_deck_path = _save_uploaded_file(previous_deck, tmpdir, 'previous_deck.pptx')
		snapshot_path = _save_uploaded_file(snapshot, tmpdir, 'snapshot.json') if snapshot else ''
		ntp_comments_path = _save_uploaded_file(ntp_comments, tmpdir, 'ntp_comments.xlsx') if ntp_comments else ''

		# Fall back to Supabase if tracker or snapshot not uploaded
		tracker_rows = None
		prev_snapshot_data = None
		if (not tracker or not snapshot) and supabase_client:
			try:
				if not tracker:
					latest = supabase_client.table('tracker_snapshot').select('*').order('uploaded_at', desc=True).limit(1).single().execute()
					if latest.data:
						tracker_rows = json.loads(latest.data['data'])
				if not snapshot:
					previous = supabase_client.table('tracker_snapshot').select('*').order('uploaded_at', desc=True).range(1, 1).execute()
					if previous.data:
						prev_snapshot_data = {
							'rows': json.loads(previous.data[0]['data']),
							'uploaded_at': previous.data[0]['uploaded_at']
						}
			except Exception as e:
				print(f'Supabase fetch error: {e}')

		if not tracker_path and tracker_rows is None:
			return jsonify({'error': 'Missing tracker: upload a .xlsx file or ensure Supabase has a snapshot'}), 400

		# Normalize deck_date to expected format if possible (MM/DD/YYYY)
		try:
			parsed = datetime.strptime(deck_date, '%m/%d/%Y')
			deck_date_str = parsed.strftime('%m/%d/%Y')
		except Exception:
			deck_date_str = deck_date

		# Call build_deck.build
		out = build_deck.build(
			tracker_path=tracker_path,
			previous_deck_path=previous_deck_path,
			snapshot_path=snapshot_path,
			ntp_comments_path=ntp_comments_path,
			deck_date=deck_date_str,
			output_dir=tmpdir,
			tracker_rows=tracker_rows,
			prev_snapshot_data=prev_snapshot_data,
		)

		# Create ZIP with outputs
		zip_buffer = io.BytesIO()
		with zipfile.ZipFile(zip_buffer, 'w', zipfile.ZIP_DEFLATED) as z:
			for key in ('deck_path', 'snapshot_path', 'ntp_comments_path', 'debug_slides_path'):
				path = out.get(key)
				if path and os.path.exists(path):
					z.write(path, arcname=os.path.basename(path))

		zip_buffer.seek(0)
		filename = f'ciege_outputs_{deck_date_str.replace("/","-")}.zip'
		resp = send_file(
			zip_buffer,
			mimetype='application/zip',
			as_attachment=True,
			download_name=filename,
		)
		if out.get('ntp_emails'):
			encoded = base64.b64encode(json.dumps(out['ntp_emails']).encode()).decode()
			resp.headers['X-Ntp-Emails'] = encoded
			resp.headers['Access-Control-Expose-Headers'] = 'X-Ntp-Emails'
		return resp

	except Exception as e:
		tb = traceback.format_exc()
		return jsonify({'error': str(e), 'traceback': tb}), 500

	finally:
		try:
			shutil.rmtree(tmpdir)
		except Exception:
			pass


V2_LATEST_ID = 'v2-build-latest'
V2_PRIOR_ID = 'v2-build-prior'


def save_v2_build(deck_bytes: bytes, filename: str, deck_date: str):
	"""Keeps the newest V2 build, rotating the previous latest into the prior slot."""
	if not supabase_client:
		return
	try:
		old = supabase_client.table('report_snapshots').select('*').eq('id', V2_LATEST_ID).execute()
		if old.data:
			supabase_client.table('report_snapshots').upsert({**old.data[0], 'id': V2_PRIOR_ID}).execute()
		supabase_client.table('report_snapshots').upsert({
			'id': V2_LATEST_ID,
			'filename': filename,
			'uploaded_at': datetime.utcnow().isoformat(),
			'data': json.dumps({'deck_date': deck_date, 'b64': base64.b64encode(deck_bytes).decode()}),
		}).execute()
	except Exception as e:
		print(f'save_v2_build failed: {e}')


def _slide_summary(prs):
	slides = []
	for idx, slide in enumerate(prs.slides):
		texts, tables, charts = [], [], []
		for sh in slide.shapes:
			if sh.has_text_frame and sh.text_frame.text.strip():
				texts.append(sh.text_frame.text.strip())
			if sh.has_table:
				rows = [[sh.table.cell(r, c).text.strip() for c in range(len(sh.table.columns))] for r in range(len(sh.table.rows))]
				tables.append(rows)
			if sh.has_chart:
				try:
					ch = sh.chart
					cats = [str(c) for c in ch.plots[0].categories]
					series = [{'name': s.name, 'values': [None if v is None else float(v) for v in s.values]} for s in ch.series]
					charts.append({'categories': cats, 'series': series})
				except KeyError:
					texts.append('[chart missing from this build]')
		slides.append({'index': idx, 'texts': texts, 'tables': tables, 'charts': charts})
	return slides


@app.route('/v2_slides', methods=['GET'])
def v2_slides_endpoint():
	which = request.args.get('which', 'latest')
	row_id = V2_PRIOR_ID if which == 'prior' else V2_LATEST_ID
	if not supabase_client:
		return jsonify({'error': 'Supabase not configured'}), 500
	res = supabase_client.table('report_snapshots').select('*').eq('id', row_id).execute()
	if not res.data:
		return jsonify({'error': f'No {which} build saved yet'}), 404
	row = res.data[0]
	payload = json.loads(row['data'])
	prs = Presentation(io.BytesIO(base64.b64decode(payload['b64'])))
	return jsonify({
		'which': which,
		'filename': row['filename'],
		'uploaded_at': row['uploaded_at'],
		'deck_date': payload.get('deck_date'),
		'slides': _slide_summary(prs),
	})


@app.route('/build_v2', methods=['POST'])
def build_v2_endpoint():
	"""Deck Builder V2 — fully separate from /build above (which it never
	calls into, and never modifies any state /build also reads). See
	build_deck_v2.py's module docstring for the full list of behavioral
	differences this exists to deliver."""
	tmpdir = tempfile.mkdtemp(prefix='ciege_build_v2_')
	try:
		tracker = request.files.get('tracker')
		previous_deck = request.files.get('previous_deck')
		decom_pptx = request.files.get('decom_pptx')
		scop_pptx = request.files.get('scop_pptx')
		deck_date = request.form.get('deck_date') or request.args.get('deck_date')

		if not previous_deck or not deck_date:
			return jsonify({'error': 'Missing required fields: previous_deck, deck_date'}), 400
		if not decom_pptx or not scop_pptx:
			return jsonify({'error': 'Missing required fields: decom_pptx, scop_pptx (generated client-side before upload)'}), 400

		tracker_path = _save_uploaded_file(tracker, tmpdir, 'tracker.xlsx') if tracker else ''
		previous_deck_path = _save_uploaded_file(previous_deck, tmpdir, 'previous_deck.pptx')
		decom_pptx_bytes = decom_pptx.read()
		scop_pptx_bytes = scop_pptx.read()

		tracker_rows = None
		if not tracker and supabase_client:
			try:
				latest = supabase_client.table('tracker_snapshot').select('*').order('uploaded_at', desc=True).limit(1).single().execute()
				if latest.data:
					tracker_rows = json.loads(latest.data['data'])
			except Exception as e:
				print(f'Supabase tracker fetch error: {e}')

		if not tracker_path and tracker_rows is None:
			return jsonify({'error': 'Missing tracker: upload a .xlsx file or ensure Supabase has a snapshot'}), 400

		try:
			parsed = datetime.strptime(deck_date, '%m/%d/%Y')
			deck_date_str = parsed.strftime('%m/%d/%Y')
		except Exception:
			deck_date_str = deck_date

		out = build_deck_v2.build(
			tracker_path=tracker_path,
			previous_deck_path=previous_deck_path,
			deck_date=deck_date_str,
			output_dir=tmpdir,
			tracker_rows=tracker_rows,
			supabase_client=supabase_client,
			decom_pptx_bytes=decom_pptx_bytes,
			scop_pptx_bytes=scop_pptx_bytes,
		)

		# Read into memory before the `finally` block's tmpdir cleanup runs —
		# send_file() given a path can still be mid-stream when that cleanup
		# fires, since Flask sends the body after this function returns (the
		# existing /build route sidesteps this the same way, via an in-memory
		# zip buffer instead of a path).
		with open(out['deck_path'], 'rb') as f:
			deck_bytes = f.read()
		save_v2_build(deck_bytes, os.path.basename(out['deck_path']), deck_date_str)
		buffer = io.BytesIO(deck_bytes)
		return send_file(
			buffer,
			mimetype='application/vnd.openxmlformats-officedocument.presentationml.presentation',
			as_attachment=True,
			download_name=os.path.basename(out['deck_path']),
		)

	except Exception as e:
		tb = traceback.format_exc()
		return jsonify({'error': str(e), 'traceback': tb}), 500

	finally:
		try:
			shutil.rmtree(tmpdir)
		except Exception:
			pass


@app.route('/ntp_emails', methods=['POST'])
def ntp_emails_endpoint():
	try:
		ntp_file = request.files.get('ntp_comments')
		if not ntp_file:
			return jsonify({'error': 'No NTP Comments file uploaded'}), 400
		with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as tmp:
			ntp_file.save(tmp.name)
			ntp_path = tmp.name
		emails = build_deck.generate_ntp_emails_from_file(ntp_path)
		os.unlink(ntp_path)
		return jsonify(emails)
	except Exception as e:
		tb = traceback.format_exc()
		return jsonify({'error': str(e), 'traceback': tb}), 500


@app.route('/gr_data', methods=['POST'])
def gr_data_endpoint():
	try:
		if not supabase_client:
			return jsonify({'error': 'Supabase not configured'}), 500
		body = request.get_json(silent=True) or {}
		gc_filter = body.get('gc') or None
		groups = gr_tracker.load_gr_data(supabase_client, gc_filter=gc_filter)
		return jsonify(groups)
	except Exception as e:
		tb = traceback.format_exc()
		return jsonify({'error': str(e), 'traceback': tb}), 500


@app.route('/ai_assistant', methods=['POST'])
def ai_assistant_endpoint():
	try:
		app.logger.info('ai_assistant endpoint hit')

		if not ANTHROPIC_API_KEY:
			app.logger.error('ai_assistant error: ANTHROPIC_API_KEY not configured on server')
			return jsonify({'error': 'ANTHROPIC_API_KEY not configured on server'}), 500

		body = request.get_json(silent=True) or {}
		messages = body.get('messages') or []
		context = body.get('context') or ''
		app.logger.info(f'Request data keys: {list(body.keys())}, message count: {len(messages)}, context length: {len(context)}')

		if not messages:
			return jsonify({'error': 'No messages provided'}), 400

		system_prompt = f"{AI_ASSISTANT_SYSTEM_PROMPT}\n\n{context}"

		anthropic_resp = requests.post(
			'https://api.anthropic.com/v1/messages',
			headers={
				'x-api-key': ANTHROPIC_API_KEY,
				'anthropic-version': '2023-06-01',
				'content-type': 'application/json',
			},
			json={
				'model': ANTHROPIC_MODEL,
				'max_tokens': 1000,
				'system': system_prompt,
				'messages': messages,
			},
			timeout=60,
		)

		if anthropic_resp.status_code != 200:
			# This is the most likely source of the 502s reported against this
			# endpoint — the Flask route itself returns 502 whenever Anthropic's
			# API responds with anything other than 200 (bad/missing API key,
			# invalid model id, rate limit, malformed request body, etc.).
			# Logged here (in addition to the JSON body already returned) since
			# that JSON body reaching Railway's logs is what the caller needs
			# to actually see the upstream response text.
			app.logger.error(
				f'ai_assistant error: Anthropic API returned {anthropic_resp.status_code} — {anthropic_resp.text[:500]}'
			)
			return jsonify({
				'error': f'Anthropic API error: {anthropic_resp.status_code}',
				'detail': anthropic_resp.text[:500],
			}), 502

		data = anthropic_resp.json()
		answer = ''.join(
			block.get('text', '') for block in data.get('content', []) if block.get('type') == 'text'
		).strip()

		return jsonify({'response': answer})
	except Exception as e:
		tb = traceback.format_exc()
		app.logger.error(f'ai_assistant error: {str(e)}')
		app.logger.error(tb)
		return jsonify({'error': str(e), 'traceback': tb}), 500


# ─────────────────────────────────────────────
# NTP comments (Deck Builder V2) — the NTP tab reads these, edits save here,
# and the Excel export is generated on demand for the customer.
# ─────────────────────────────────────────────

def _ntp_months(deck_date_str: str):
	"""Pending-NTP HOPs for every month in the window around the call, with the
	saved comment and status merged in. Same POR logic as the V2 slides."""
	import pandas as pd
	if not supabase_client:
		raise RuntimeError('Supabase is not configured')
	latest = supabase_client.table('tracker_snapshot').select('*').order('uploaded_at', desc=True).limit(1).single().execute()
	if not latest.data:
		raise RuntimeError('No tracker snapshot in Supabase')
	tracker_rows = json.loads(latest.data['data'])
	call_dt = build_deck_v2.call_tuesday(deck_date_str)
	data = build_deck.extract_data('', '', '', call_dt.strftime('%m/%d/%Y'), tracker_rows=tracker_rows, prev_snapshot_data=None)
	# Start three months back so Jul–Sep show up even though the call is in Oct.
	por, order = build_deck_v2.compute_por_window(data['df'], pd.Timestamp(call_dt - timedelta(days=90)), window_months=12, max_active=12)
	store = ntp_comments_v2.load_store(supabase_client)
	months = []
	for mo_key, mo_name, sheet in order:
		rows = ntp_comments_v2.month_rows(por[mo_key]['pending_rows'], store.get(sheet, {}))
		months.append({'sheet': sheet, 'label': mo_name, 'rows': rows})
	return call_dt, months


@app.route('/ntp_comments', methods=['GET'])
def ntp_comments_get():
	try:
		deck_date = request.args.get('deck_date') or datetime.now().strftime('%m/%d/%Y')
		call_dt, months = _ntp_months(deck_date)
		return jsonify({
			'deck_date': call_dt.strftime('%m/%d/%Y'),
			'statuses': ntp_comments_v2.STATUS_OPTIONS,
			'months': months,
		})
	except Exception as e:
		traceback.print_exc()
		return jsonify({'error': str(e)}), 500


@app.route('/ntp_comments', methods=['POST'])
def ntp_comments_save():
	try:
		body = request.get_json(force=True) or {}
		sheet, hop = body.get('sheet'), body.get('hop')
		if not sheet or not hop:
			return jsonify({'error': 'sheet and hop are required'}), 400
		if 'status' in body and body['status'] not in ntp_comments_v2.STATUS_OPTIONS:
			return jsonify({'error': 'unknown status'}), 400
		store = ntp_comments_v2.load_store(supabase_client)
		ntp_comments_v2.set_entry(store, sheet, hop, comment=body.get('comment'), status=body.get('status'))
		ntp_comments_v2.save_store(supabase_client, store)
		return jsonify({'ok': True})
	except Exception as e:
		traceback.print_exc()
		return jsonify({'error': str(e)}), 500


@app.route('/ntp_comments/export', methods=['GET'])
def ntp_comments_export():
	try:
		deck_date = request.args.get('deck_date') or datetime.now().strftime('%m/%d/%Y')
		call_dt, months = _ntp_months(deck_date)
		xlsx = ntp_comments_v2.export_workbook([(m['sheet'], m['rows']) for m in months])
		return send_file(io.BytesIO(xlsx), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
		                 as_attachment=True, download_name=f'NTP_Pending_{call_dt.strftime("%m-%d-%Y")}.xlsx')
	except Exception as e:
		traceback.print_exc()
		return jsonify({'error': str(e)}), 500


if __name__ == '__main__':
	port = int(os.environ.get('PORT', 8000))
	app.run(host='0.0.0.0', port=port, debug=False)
